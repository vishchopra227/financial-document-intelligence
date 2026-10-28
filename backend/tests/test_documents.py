"""Automated tests for Document Ingestion and Management API."""

import io
from pathlib import Path
import pytest
import pypdf
import docx

from backend.app.config import settings


def create_sample_pdf() -> bytes:
    """Create a minimal valid multi-page PDF in memory using pypdf."""
    writer = pypdf.PdfWriter()
    
    # Page 1
    page1 = writer.add_blank_page(width=612, height=792)
    # Page 2
    page2 = writer.add_blank_page(width=612, height=792)
    
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def create_sample_docx() -> bytes:
    """Create a valid docx with paragraphs and a table in memory."""
    doc = docx.Document()
    doc.add_heading("Apple Inc. Form 10-K Excerpt", level=1)
    doc.add_paragraph("Total revenue for fiscal year 2023 was $383,285 million.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "FY 2023"
    table.cell(1, 0).text = "Net Sales"
    table.cell(1, 1).text = "$383.28B"
    
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_upload_pdf_document(client):
    pdf_bytes = create_sample_pdf()
    files = {"file": ("sample_report.pdf", pdf_bytes, "application/pdf")}
    data = {"workspace_id": "test_ws"}

    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 201
    doc_data = response.json()
    assert doc_data["original_filename"] == "sample_report.pdf"
    assert doc_data["file_type"] == "pdf"
    assert doc_data["page_count"] == 2
    assert doc_data["status"] in ("extracted", "indexed")
    assert doc_data["workspace_id"] == "test_ws"

    doc_id = doc_data["id"]

    # Verify physical file stored safely
    storage_path = settings.UPLOAD_DIR / f"{doc_id}.pdf"
    assert storage_path.exists()

    # Cleanup
    del_resp = client.delete(f"/api/documents/{doc_id}")
    assert del_resp.status_code == 200
    assert not storage_path.exists()


def test_upload_txt_document(client):
    txt_content = (
        "FINANCIAL REPORT 2023\n"
        "Revenue: $54.2 Billion\n"
        "Operating Expenses: $12.1 Billion\n"
        "Net Profit Margin: 21.4%\n"
    ).encode("utf-8")

    files = {"file": ("earnings_release.txt", txt_content, "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 201
    data = response.json()
    assert data["file_type"] == "txt"
    assert data["page_count"] >= 1
    assert "FINANCIAL REPORT 2023" in data["document_summary"]

    doc_id = data["id"]
    # Check status endpoint
    status_resp = client.get(f"/api/documents/{doc_id}/status")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["status"] in ("extracted", "indexed")

    # Check preview endpoint
    preview_resp = client.get(f"/api/documents/{doc_id}/preview")
    assert preview_resp.status_code == 200
    preview_data = preview_resp.json()
    assert len(preview_data["pages"]) >= 1
    assert "Revenue: $54.2 Billion" in preview_data["pages"][0]["text"]

    # Cleanup
    client.delete(f"/api/documents/{doc_id}")


def test_upload_docx_document(client):
    docx_bytes = create_sample_docx()
    files = {"file": ("report.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 201
    data = response.json()
    assert data["file_type"] == "docx"
    assert data["page_count"] >= 1

    doc_id = data["id"]
    preview_resp = client.get(f"/api/documents/{doc_id}/preview")
    assert preview_resp.status_code == 200
    preview_text = preview_resp.json()["pages"][0]["text"]
    assert "Total revenue for fiscal year 2023" in preview_text
    assert "[TABLE START]" in preview_text
    assert "Net Sales | $383.28B" in preview_text

    # Cleanup
    client.delete(f"/api/documents/{doc_id}")


def test_reject_invalid_file_type(client):
    dummy_exe = b"MZ\x90\x00\x03\x00\x00\x00"
    files = {"file": ("malware.exe", dummy_exe, "application/octet-stream")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "Unsupported file format" in response.json()["detail"]


def test_reject_empty_file(client):
    files = {"file": ("empty.txt", b"", "text/plain")}
    response = client.post("/api/documents/upload", files=files)
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_list_and_search_documents(client):
    txt_content = b"Financial Overview Company ABC"
    files = {"file": ("alphabet_annual_report.txt", txt_content, "text/plain")}
    res = client.post("/api/documents/upload", files=files)
    doc_id = res.json()["id"]

    # List all
    list_res = client.get("/api/documents")
    assert list_res.status_code == 200
    list_data = list_res.json()
    assert list_data["total_count"] >= 1
    filenames = [d["original_filename"] for d in list_data["documents"]]
    assert "alphabet_annual_report.txt" in filenames

    # Search filter
    search_res = client.get("/api/documents?search=alphabet")
    assert search_res.status_code == 200
    search_data = search_res.json()
    assert len(search_data["documents"]) == 1
    assert search_data["documents"][0]["id"] == doc_id

    # Non-matching search
    no_match_res = client.get("/api/documents?search=nonexistent_xyz")
    assert no_match_res.status_code == 200
    assert len(no_match_res.json()["documents"]) == 0

    # Cleanup
    client.delete(f"/api/documents/{doc_id}")


def test_get_nonexistent_document(client):
    res = client.get("/api/documents/nonexistent-id-000")
    assert res.status_code == 404
